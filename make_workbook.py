import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

data = json.load(open('hiv_dashboard_data.json'))
meta = data['meta']
entries = data['entries']

ROUTES = ['PO','SC','IM','IV','VR','TD']
FREQ_MAP={'Q1W':'1W','Q2W':'2W','Q1M':'1M','Q2M':'2M','Q3M':'3M','Q4M':'4M','Q6M':'6M','Q12M':'12M'}
ROUTE_MAP={'PO':'PO','SC':'SC','IM':'IM','IV':'IV','Vaginal ring':'VR','Transdermal':'TD'}
FREQS  = ['1W','2W','1M','2M','3M','4M','6M','12M']
STAGES = ['Approved','Late-stage','Early-stage','Not stated']
PHASES = ['Preclinical','Phase I','Phase I/II','Phase II','Phase II/III','Phase III','Phase IV']
BANDS  = ['formulations','compounds']
INDIC  = ['Treatment','Prevention','Both']
FREQSRC= ['LAPaL record','derived from trials','manually confirmed','not stated']

# ---- styling ----
INK='12161C'; BLUE='153F8F'; MAG='A21C7A'; SOFT='5A6472'
HEAD_FILL=PatternFill('solid', fgColor=BLUE)
GRP_ROUTE=PatternFill('solid', fgColor='E8EEF7')
GRP_FREQ =PatternFill('solid', fgColor='F3E8F0')
EDIT_FILL=PatternFill('solid', fgColor='FFFFFF')
LOCK_FILL=PatternFill('solid', fgColor='F0EFEC')
thin=Side(style='thin', color='DFE3E8')
BORD=Border(left=thin,right=thin,top=thin,bottom=thin)
HFONT=Font(name='Arial', bold=True, color='FFFFFF', size=9)
CFONT=Font(name='Arial', size=9, color=INK)
NFONT=Font(name='Arial', size=9, bold=True, color=INK)

wb = openpyxl.Workbook()

# ============================================================ ENTRIES
ws = wb.active
ws.title = 'Entries'

# column plan: (header, key, group)
cols = [('id','id','lock'),('band','band','base'),('name','name','base'),
        ('indication','indication','base'),('stage','stage','base'),
        ('highest_phase','highest_phase','base'),('developers','developers','base'),
        ('drug_class','drug_class','base')]
cols += [(r,('route',r),'route') for r in ROUTES]
cols += [(f,('freq',f),'freq') for f in FREQS]
cols += [('frequency_source','frequency_source','base'),
         ('exclude','exclude','base'),
         ('exclude_reason','exclude_reason','base'),
         ('flag_no_trials','flag_no_trials','base'),
         ('flag_no_reg_rows','flag_no_reg_rows','base'),
         ('curator_notes','curator_notes','base'),
         ('link in LAPaL','link_in_lapal','base')]

HROW=2
for ci,(hdr,key,grp) in enumerate(cols, start=1):
    c=ws.cell(row=HROW, column=ci, value=hdr)
    c.font=HFONT; c.alignment=Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border=BORD
    if grp=='route': c.fill=PatternFill('solid', fgColor='2E5AA8')
    elif grp=='freq': c.fill=PatternFill('solid', fgColor='7A3F6E')
    else: c.fill=HEAD_FILL

# group banner row (row 1)
def banner(c0,c1,text,fill):
    ws.merge_cells(start_row=1,start_column=c0,end_row=1,end_column=c1)
    cell=ws.cell(row=1,column=c0,value=text)
    cell.fill=fill; cell.font=Font(name='Arial',bold=True,size=8,color=INK)
    cell.alignment=Alignment(horizontal='center')
r0=[i+1 for i,(_,_,g) in enumerate(cols) if g=='route']
f0=[i+1 for i,(_,_,g) in enumerate(cols) if g=='freq']
banner(r0[0],r0[-1],'ROUTE  (mark x)',GRP_ROUTE)
banner(f0[0],f0[-1],'DOSING INTERVAL  (mark x)',GRP_FREQ)

def ind_of(e):
    t,p=e['is_treatment'],e['is_prevention']
    return 'Both' if t and p else ('Treatment' if t else 'Prevention')

r=HROW+1
for e in entries:
    rowmap={
        'id':e['id'],'band':e['band'],'name':e['name_full'],'indication':ind_of(e),
        'stage':e['stage'],'highest_phase':e['highest_phase'] or '',
        'developers':'; '.join(e['developers_full']),'drug_class':e['drug_class'] or '',
        'frequency_source':e['frequency_provenance'],
        'flag_no_trials':'x' if e['flags']['no_linked_trials'] else '',
        'flag_no_reg_rows':'x' if e['flags']['approved_but_no_regulatory_rows'] else '',
        'curator_notes':'',
    }
    for ci,(hdr,key,grp) in enumerate(cols, start=1):
        if grp=='route':
            v='x' if key[1] in [ROUTE_MAP.get(x,x) for x in e['routes']] else ''
        elif grp=='freq':
            v='x' if key[1] in [FREQ_MAP.get(x,x) for x in e['frequencies']] else ''
        else:
            v=rowmap.get(key,'')
        cell=ws.cell(row=r,column=ci,value=v)
        cell.border=BORD
        cell.font=NFONT if hdr=='name' else CFONT
        if grp in ('route','freq'):
            cell.alignment=Alignment(horizontal='center')
            cell.fill=GRP_ROUTE if grp=='route' else GRP_FREQ
        elif hdr=='id':
            cell.fill=LOCK_FILL; cell.font=Font(name='Arial',size=8,color=SOFT)
        else:
            cell.fill=EDIT_FILL
    r+=1
LAST=r-1

# validations
def dv(formula, cells, msg):
    d=DataValidation(type='list', formula1=formula, allow_blank=True, showErrorMessage=True)
    d.error=msg; d.prompt=msg
    ws.add_data_validation(d)
    for c in cells: d.add(c)

def colcells(hdr):
    ci=[i+1 for i,(h,_,_) in enumerate(cols) if h==hdr][0]
    L=get_column_letter(ci)
    return f'{L}{HROW+1}:{L}{LAST}'

dv('"'+','.join(BANDS)+'"', [colcells('band')], 'Choose formulations or compounds')
dv('"'+','.join(INDIC)+'"', [colcells('indication')], 'Treatment, Prevention or Both')
dv('"'+','.join(STAGES)+'"', [colcells('stage')], 'Pick a development stage')
dv('"'+','.join(PHASES)+'"', [colcells('highest_phase')], 'Highest clinical phase')
dv('"'+','.join(FREQSRC)+'"', [colcells('frequency_source')], 'Where the interval came from')
xcells=[colcells(h) for h in ROUTES+FREQS+['flag_no_trials','flag_no_reg_rows','exclude']]
dv('"x"', xcells, 'Type x to mark, or leave blank')

# widths + freeze
W={'id':6,'band':13,'name':40,'indication':12,'stage':12,'highest_phase':12,
   'developers':30,'drug_class':16,'frequency_source':17,'flag_no_trials':9,
   'flag_no_reg_rows':10,'curator_notes':26,'exclude':9,'exclude_reason':34,'link in LAPaL':46}
for ci,(hdr,_,grp) in enumerate(cols, start=1):
    ws.column_dimensions[get_column_letter(ci)].width = W.get(hdr, 5)
ws.row_dimensions[1].height=15
ws.row_dimensions[2].height=34
ws.freeze_panes='D3'   # keep id/band/name + header visible

# ============================================================ LEGEND
lg=wb.create_sheet('Legend & how to use')
lg.column_dimensions['A'].width=26; lg.column_dimensions['B'].width=88
def L(row,a,b='',bold=False,size=10,color=INK,fill=None):
    ca=lg.cell(row=row,column=1,value=a); cb=lg.cell(row=row,column=2,value=b)
    ca.font=Font(name='Arial',bold=True,size=size,color=color)
    cb.font=Font(name='Arial',size=size,color=INK)
    cb.alignment=Alignment(wrap_text=True,vertical='top')
    if fill:
        for c in (ca,cb): c.fill=PatternFill('solid',fgColor=fill)
    return row+1

row=1
row=L(row,'LAPaL long-acting HIV','Curation sheet for the interactive dashboard',bold=True,size=13,color=BLUE)
row=L(row,'','')
row=L(row,'HOW IT WORKS','',fill='E8EEF7')
row=L(row,'1. Edit the Entries tab','One row per product. Change any white cell. Add or delete whole rows freely; order does not matter.')
row=L(row,'2. Save the file','Keep it as .xlsx.')
row=L(row,'3. Run the build step',"In a terminal:  python build_data.py hiv_curation.xlsx  \u2192 writes hiv_dashboard_data.json, which the dashboard reads.")
row=L(row,'4. Refresh the dashboard','The new data appears. The build step prints a check report and refuses to write if something essential is missing.')
row=L(row,'','')
row=L(row,'COLUMNS','',fill='E8EEF7')
row=L(row,'id','Leave as is. Leave blank for a new row and the build step assigns one.')
row=L(row,'band','formulations = a discrete long-acting product. compounds = an underlying molecule in LA development.')
row=L(row,'name','Product or molecule name, shown on the dashboard.')
row=L(row,'indication','Treatment, Prevention, or Both. Drives which selection the row appears under. "Both" shows it once, tagged for both.')
row=L(row,'stage','Approved / Late-stage / Early-stage / Not stated. Sets the row colour and grouping.')
row=L(row,'highest_phase','Highest clinical phase on record. Optional, shown in the tooltip.')
row=L(row,'developers','One or more, separated by a semicolon ( ; ).')
row=L(row,'drug_class','Optional, shown in the tooltip.')
row=L(row,'ROUTE columns','Mark x under every route that applies. PO oral, SC subcutaneous, IM intramuscular, IV intravenous, VR vaginal ring, TD transdermal.')
row=L(row,'INTERVAL columns','Mark x under every dosing interval studied. 1W weekly, 2W 2-weekly, 1M monthly, 2M, 3M, 4M, 6M, 12M yearly. Leave all blank if unknown; the row then sits in the "not stated" lane. Do not enter daily or single-dose.')
row=L(row,'frequency_source','LAPaL record = confirmed. "derived from trials" = filled from a linked trial and shown with a dagger until you confirm. Change to "manually confirmed" once checked. "not stated" = no interval recorded at all; use it on rows where every interval column is blank.')
row=L(row,'flag_no_trials','Mark x if no clinical trials are linked in LAPaL. Surfaced in the data-quality panel.')
row=L(row,'flag_no_reg_rows','Mark x if approved but missing structured regulatory rows. Surfaced in the data-quality panel.')
row=L(row,'exclude','Leave blank to show the entry. Mark x to withhold it from the dashboard. The row is still validated, and the dashboard reports how many entries were withheld, so nothing disappears silently.')
row=L(row,'exclude_reason','Why the entry is withheld. Shown in the data-quality panel next to the entry name.')
row=L(row,'curator_notes','Free text for you. Not shown on the dashboard.')
row=L(row,'link in LAPaL','Full web address of this entry on lapal.ch. Clicking the row in the dashboard opens it in a new tab. Must start with https://.')
row=L(row,'','')
row=L(row,'AUTOMATIC FLAGS','The build step computes these; you do not enter them:',fill='E8EEF7')
row=L(row,'missing_frequency','set when no interval x is marked.')
row=L(row,'conventional_approval_only','set when stage is Approved, no interval is marked, and the only route is oral (an approved conventional drug shown as an API in LA development).')
row=L(row,'','')
row=L(row,'RULES THE BUILD STEP CHECKS','',fill='F3E8F0')
row=L(row,'every row needs','a band, a name, an indication and a stage. Missing any of these stops the build with a clear message.')
row=L(row,'controlled values','band, stage, indication and interval codes must match the lists above. Typos are reported by row.')

# example row note on Entries
note=ws.cell(row=LAST+3, column=26,
    value='Example: to add a product, copy a row, change name, set band/indication/stage, '
          'mark x under each route and interval, leave id blank. Then run build_data.py.')
note.font=Font(name='Arial', italic=True, size=8, color=SOFT)

wb.save('hiv_curation.xlsx')
print('wrote hiv_curation.xlsx  rows:', LAST-HROW)
